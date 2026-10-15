#!/usr/bin/env python3
"""HTML chart browser for sheet 1. Site List / 2. HourlyChartOfIssueSites.

Reads a published workbook and writes one HTML file. The Excel report is not modified.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

EXCEL_EPOCH = datetime(1899, 12, 30)
HTML_VERSION = "22.6"
EXCEL_VERSION = "v22"
FLAG_NAMES = (
    "DL Throughput<7Mbps",
    "Data Vol>50GB",
    "DL Throughput<5Mbps",
    "Data Vol>100GB",
)


def _num(value) -> float | None:
    if isinstance(value, (int, float)) and value == value:
        return float(value)
    return None


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_date(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(value, (int, float)) and value > 1000:
        return EXCEL_EPOCH + timedelta(days=float(value))
    return None


def _as_hour(value) -> int | None:
    if isinstance(value, datetime):
        return int(value.hour)
    if isinstance(value, time):
        return int(value.hour)
    if isinstance(value, (int, float)):
        number = float(value)
        if 0 <= number <= 23 and number == int(number):
            return int(number)
        if 0 <= number < 1:
            return int(round(number * 24)) % 24
    text = _text(value)
    if len(text) >= 2 and text[:2].isdigit():
        return int(text[:2])
    return None


def _round(value: float | None, digits: int = 2):
    if value is None:
        return None
    return round(value, digits)


def load_chart_view(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    sites_ws = book["1. Site List"]
    period = ""
    sites = []
    for row in sites_ws.iter_rows(min_row=1, max_col=18, values_only=True):
        if row[0] == "No.":
            continue
        if isinstance(row[0], str) and not period and "Oct" in row[0]:
            period = row[0].split(".", 1)[0].strip()
        if not isinstance(row[0], (int, float)):
            continue
        license_mbps = _num(row[16])
        sites.append(
            {
                "n": int(row[0]),
                "site": _text(row[1]),
                "tech": _text(row[2]),
                "siteType": _text(row[3]),
                "category": _text(row[4]),
                "severity": _text(row[5]),
                "flags": [_text(row[6]), _text(row[7]), _text(row[8]), _text(row[9])],
                "peak": _round(_num(row[10])),
                "level": _round(_num(row[11])),
                "touchHours": _num(row[12]),
                "totalHours": _num(row[13]),
                "ratio": _round(_num(row[14]), 1),
                "where": _text(row[15]),
                "license": _round(license_mbps),
                "threshold": _round(license_mbps * 0.85) if license_mbps else None,
                "gap": _round(_num(row[17]), 1),
                "listed": True,
            }
        )
    by_site = {item["site"]: item for item in sites}
    hourly = book["3. Hourly KPI"]
    hours: set[tuple[datetime, int]] = set()
    series: dict[str, dict[tuple[datetime, int], float | None]] = {item["site"]: {} for item in sites}
    for row in hourly.iter_rows(min_row=1, max_col=8, values_only=True):
        site = _text(row[2])
        if site not in by_site:
            continue
        day = _as_date(row[0])
        hour = _as_hour(row[1])
        if day is None or hour is None:
            continue
        key = (day, hour)
        hours.add(key)
        series[site][key] = _round(_num(row[3]))
    book.close()
    axis = sorted(hours)
    labels = [f"{day.day}/{day.strftime('%b')} {hour:02d}:00" for day, hour in axis]
    for item in sites:
        values = series[item["site"]]
        item["rx"] = [values.get(key) for key in axis]
    return {"period": period, "labels": labels, "flagNames": list(FLAG_NAMES), "sites": sites}


OSS_1H_NAMES = (
    "DL Throughput 1h (Mbps)",
    "DLPRBUtilization,%",
    "Data Volume 1h (GB)",
    "Max User 1h",
    "Mean User 1h",
)


def load_oss_1h(path: Path) -> dict[str, list]:
    """1 BusyHour OSS KPIs from URGENTBWINC_BusyHour, keyed by the 7-character site code.

    Each value is [data block, user block]. A block is throughput, PRB, volume, max user, mean user.
    """
    book = load_workbook(path, read_only=True, data_only=True)
    ws = book["URGENTBWINC_BusyHour"]
    header = None
    site_col = 1
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        if row and "DL Throughput 1h (Mbps)" in row:
            header = list(row)
            break
    if header is None:
        book.close()
        raise SystemExit("URGENTBWINC_BusyHour has no 1h throughput header")
    throughput_cols = [i for i, name in enumerate(header) if name == "DL Throughput 1h (Mbps)"]
    if len(throughput_cols) < 2:
        book.close()
        raise SystemExit("Expected Data and User 1h throughput columns")
    blocks = []
    for start in throughput_cols[:2]:
        names = header[start:start + len(OSS_1H_NAMES)]
        if tuple(names) != OSS_1H_NAMES:
            book.close()
            raise SystemExit(f"1h KPI columns changed: {names}")
        blocks.append(start)
    if "eNodeB Name" in header:
        site_col = header.index("eNodeB Name")
    oss: dict[str, list] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        site = _text(row[site_col]) if row and len(row) > site_col else ""
        if not site or site == "eNodeB Name":
            continue
        pair = []
        for start in blocks:
            pair.append([_round(_num(row[start + offset])) for offset in range(len(OSS_1H_NAMES))])
        oss[site[:7]] = pair
    book.close()
    return oss


def _axis_keys(labels: list[str], year: int) -> list[tuple]:
    keys = []
    for label in labels:
        day_text, clock = label.split()
        day, month_name = day_text.split("/")
        month = datetime.strptime(month_name, "%b").month
        keys.append((datetime(year, month, int(day)).date(), int(clock[:2])))
    return keys


def add_all_sites(view: dict, feg_path: Path, license_path: Path, geo_path: Path) -> None:
    """Append every ERBS from the FEG window. Issue sites already on the list stay as they are."""
    import sys

    import numpy as np
    import pandas as pd

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import fege_choke_check as fcc
    import fege_choke_check_ericsson as eric

    year = 2026
    for token in view["period"].replace("–", " ").replace("-", " ").split():
        if token.isdigit() and len(token) == 4:
            year = int(token)
    axis = _axis_keys(view["labels"], year)
    axis_pos = {key: index for index, key in enumerate(axis)}
    known = {item["site"] for item in view["sites"]}

    raw = eric.load_ericsson(feg_path)
    frame = pd.DataFrame(
        {
            "site": raw[eric.ERIC_SITE_COL].astype(str).str.strip(),
            "date": pd.to_datetime(fcc.parse_kpi_date(raw["Date"])).dt.date,
            "hour": raw["Hour"].astype(int),
            "rx": pd.to_numeric(raw[eric.ERIC_DL_COL], errors="coerce"),
        }
    )
    frame = frame[frame["date"].isin({day for day, _hour in axis})]
    frame = frame.drop_duplicates(["site", "date", "hour"], keep="last")

    license_of: dict[str, float] = {}
    if license_path.exists():
        lic = pd.read_excel(license_path, engine="pyxlsb")
        caps = pd.to_numeric(lic["dlBbCapacityNet"], errors="coerce")
        license_of = {
            str(node).strip(): float(cap)
            for node, cap in zip(lic["NodeId"], caps)
            if pd.notna(cap) and float(cap) > 0
        }
    tech_of: dict[str, str] = {}
    type_of: dict[str, str] = {}
    if geo_path.exists():
        geo = fcc.load_site_geo(geo_path)
        tech_of = fcc.tech_by_site_from_geo(geo)
        type_of = fcc.site_type_by_site_from_geo(geo)

    added = 0
    for site, group in frame.groupby("site", sort=True):
        if site in known or not site:
            continue
        series = [None] * len(axis)
        hours = group["hour"].to_numpy()
        values = group["rx"].to_numpy(dtype=float)
        days = group["date"].to_numpy()
        finite_mask = np.isfinite(values)
        for day, hour, value, ok in zip(days, hours, values, finite_mask):
            pos = axis_pos.get((day, int(hour)))
            if pos is None:
                continue
            series[pos] = round(float(value), 2) if ok else None
        finite = values[finite_mask]
        peak = float(np.max(finite)) if len(finite) else None
        level = peak * 0.95 if peak is not None else None
        touch = finite_mask & (values >= level) if level is not None else finite_mask & False
        hours_n = int(finite_mask.sum())
        hours_on = int(touch.sum())
        busy = (hours >= fcc.BUSY_START) & (hours <= fcc.BUSY_END) & touch
        off = touch & ~((hours >= fcc.BUSY_START) & (hours <= fcc.BUSY_END))
        license_mbps = license_of.get(site)
        code = site[:7]
        gap = None
        if license_mbps and peak:
            gap = round((license_mbps - peak) / license_mbps * 100.0, 1)
        view["sites"].append(
            {
                "n": None,
                "site": site,
                "tech": tech_of.get(code, "4G"),
                "siteType": type_of.get(code, "Not found"),
                "category": "",
                "severity": "",
                "flags": ["", "", "", ""],
                "peak": round(peak, 2) if peak is not None else None,
                "level": round(level, 2) if level is not None else None,
                "touchHours": hours_on,
                "totalHours": hours_n,
                "ratio": round(100.0 * hours_on / hours_n, 1) if hours_n else None,
                "where": fcc.cap_window(int(busy.sum()), int(off.sum())) if hours_on else "",
                "license": round(license_mbps, 2) if license_mbps else None,
                "threshold": round(license_mbps * 0.85, 2) if license_mbps else None,
                "gap": gap,
                "rx": series,
                "listed": False,
            }
        )
        added += 1
    print(f"All Sites: {len(view['sites']):,} ERBS ({added:,} not on 1. Site List)", flush=True)


HTML = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Hourly charts — Excel __EXCEL_VERSION__</title>
<style>
  :root {
    --navy: #1F4E79;
    --ink: #1a1a1a;
    --line: #d9d9d9;
    --paper: #f4f6f8;
    --severe: #c00000;
    --high: #c65911;
    --medium: #bf8f00;
    --low: #548235;
    --decision: #1a5276;
  }
  * { box-sizing: border-box; }
  html, body { margin: 0; height: 100%; }
  body {
    font-family: Calibri, "Segoe UI", sans-serif;
    color: var(--ink);
    background: var(--paper);
    display: grid;
    grid-template-columns: 248px 1fr;
    height: 100vh;
  }
  aside {
    background: #fff;
    border-right: 1px solid var(--line);
    display: flex;
    flex-direction: column;
    min-height: 0;
  }
  header, .chart-head { padding: 14px 16px 10px; }
  aside header { padding: 6px 8px 2px; }
  .chart-head { padding: 8px 16px 4px; }
  h1 { margin: 0; font-size: 18px; color: var(--navy); }
  aside h1 { font-size: 13px; }
  .sub { margin: 4px 0 0; color: #595959; font-size: 12px; }
  aside .sub { margin-top: 1px; font-size: 10px; line-height: 1.2; }
  .tools { padding: 0 8px 4px; display: flex; flex-direction: column; gap: 4px; }
  input[type="search"] {
    width: 100%;
    border: 1px solid #b0b0b0;
    border-radius: 4px;
    padding: 3px 6px;
    font: inherit;
    font-size: 11px;
  }
  .filters { display: flex; flex-wrap: wrap; gap: 3px; }
  .filters button, .nav button {
    border: 1px solid #b0b0b0;
    background: #fff;
    border-radius: 12px;
    padding: 4px 10px;
    font: inherit;
    font-size: 12px;
    cursor: pointer;
  }
  .filters button { padding: 1px 6px; font-size: 10px; }
  .filters button.on { background: var(--navy); color: #fff; border-color: var(--navy); }
  #list { overflow: auto; flex: 1; position: relative; }
  #spacer { position: relative; }
  .site {
    position: absolute;
    left: 0;
    right: 0;
    height: 32px;
    text-align: left;
    border: 0;
    border-bottom: 1px solid #eee;
    background: #fff;
    padding: 2px 8px;
    cursor: pointer;
    font: inherit;
    overflow: hidden;
  }
  .site.on { background: #deebf7; }
  .site .name { font-weight: 700; color: var(--navy); font-size: 12px; line-height: 1.15; white-space: nowrap; }
  .site .meta { color: #595959; font-size: 10px; margin-top: 0; line-height: 1.15; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  .sev { font-weight: 700; }
  .sev.Severe { color: var(--severe); }
  .sev.High { color: var(--high); }
  .sev.Medium { color: var(--medium); }
  .sev.Low { color: var(--low); }
  .sev.Decision { color: var(--decision); }
  main { min-width: 0; min-height: 0; display: flex; flex-direction: column; }
  .chart-head { display: flex; justify-content: space-between; gap: 12px; align-items: flex-start; }
  .kpis { display: flex; flex-wrap: wrap; gap: 4px; padding: 0 16px 4px; }
  .kpi {
    background: #fff;
    border: 1px solid var(--line);
    border-radius: 3px;
    padding: 2px 6px;
  }
  .kpi b { display: block; font-size: 12px; line-height: 1.15; white-space: nowrap; }
  .kpi span { color: #595959; font-size: 9px; line-height: 1.15; white-space: nowrap; }
  .oss { display: grid; grid-template-columns: 1fr 1fr; gap: 6px; padding: 0 16px 4px; }
  .oss-block { background: #fff; border: 1px solid var(--line); border-radius: 3px; overflow: hidden; }
  .oss-block h2 { margin: 0; padding: 2px 8px; color: #fff; font-size: 11px; font-weight: 700; line-height: 1.25; }
  .oss-row { display: flex; justify-content: space-between; gap: 8px; padding: 1px 8px; font-size: 11px; line-height: 1.25; border-top: 1px solid #eee; }
  .oss-row span { color: #595959; }
  .oss-row b { font-variant-numeric: tabular-nums; }
  .oss-empty { padding: 3px 8px; color: #595959; font-size: 11px; }
  #stage { flex: 1; min-height: 280px; margin: 0 16px 8px; background: #fff; border: 1px solid #d0d0d0; position: relative; }
  canvas { width: 100%; height: 100%; display: block; }
  .legend { display: flex; flex-wrap: wrap; gap: 8px; padding: 2px 16px 2px; font-size: 11px; }
  .legend label { display: flex; align-items: center; gap: 6px; cursor: pointer; }
  .swatch { width: 22px; height: 0; border-top-width: 3px; border-top-style: solid; }
  .swatch.dash { border-top-style: dashed; }
  #tip {
    position: absolute; display: none; pointer-events: none;
    background: #fff; border: 1px solid #b0b0b0; border-radius: 4px;
    padding: 6px 8px; font-size: 12px; box-shadow: 0 2px 8px rgba(0,0,0,.12);
  }
  .nav { display: flex; gap: 6px; }
  @media (max-width: 860px) {
    .oss { grid-template-columns: 1fr; }
    body { grid-template-columns: 1fr; grid-template-rows: 32vh 1fr; }
    aside { border-right: 0; border-bottom: 1px solid var(--line); }
  }
</style>
</head>
<body>
<aside>
  <header>
    <h1>1. Site List · Excel __EXCEL_VERSION__</h1>
    <p class="sub" id="period"></p>
  </header>
  <div class="tools">
    <input id="q" type="search" placeholder="Search ERBS, tech, severity" autocomplete="off">
    <div class="filters" id="filters"></div>
  </div>
  <div id="list"></div>
</aside>
<main>
  <div class="chart-head">
    <div>
      <h1 id="title">Hourly chart</h1>
      <p class="sub" id="caption"></p>
    </div>
    <div class="nav">
      <button id="prev" type="button">Previous</button>
      <button id="next" type="button">Next</button>
    </div>
  </div>
  <div class="kpis" id="kpis"></div>
  <div class="oss" id="oss"></div>
  <div class="legend" id="legend"></div>
  <div id="stage">
    <canvas id="chart"></canvas>
    <div id="tip"></div>
  </div>
</main>
<script>
const VIEW = __DATA__;
const SERIES = [
  { key: "rx", name: "Max DL Capacity Usage (Mbps)", color: "#5B9BD5", dash: false },
  { key: "peak", name: "Peak", color: "#C00000", dash: true },
  { key: "level", name: "Below5%Peak", color: "#ED7D31", dash: true },
  { key: "license", name: "License", color: "#548235", dash: true },
  { key: "threshold", name: "85% of license", color: "#833C0C", dash: true }
];
const on = Object.fromEntries(SERIES.map(s => [s.key, true]));
let filter = "All";
let query = "";
let current = 0;
let shown = [];
const ROW_H = 32;
const filters = ["All", "Tx Flat/Chocked", "Peak Touch", "License usage", "All Sites"];
VIEW.sites.forEach((site, index) => { site.i = index; });

function primary(site) {
  return (site.category || "").split(" + ")[0];
}
function visible() {
  const q = query.trim().toLowerCase();
  const rows = VIEW.sites.filter(site => {
    if (filter === "All Sites") {
      /* every ERBS in the FEG window */
    } else if (filter === "All") {
      if (!site.listed) return false;
    } else if (!site.listed || primary(site) !== filter) return false;
    if (!q) return true;
    return [site.site, site.tech, site.siteType, site.category, site.severity, String(site.n || "")]
      .join(" ").toLowerCase().includes(q);
  });
  if (filter === "All Sites") rows.sort((a, b) => a.site.localeCompare(b.site));
  return rows;
}
function rowTitle(site) {
  return site.listed && site.n ? site.n + ". " + site.site : site.site;
}
function rowMeta(site) {
  if (!site.listed) return "Not listed · " + (site.tech || "") + " · " + (site.siteType || "");
  return primary(site) + ' · <span class="sev ' + site.severity + '">' + site.severity + "</span> · " + site.tech + " · " + site.siteType;
}
function yScale(values) {
  const finite = values.filter(v => v != null && isFinite(v));
  const peak = finite.length ? Math.max(...finite) : 50;
  const target = Math.max(peak * 1.15, peak + 4);
  const brackets = [[25,5],[50,10],[75,25],[100,25],[125,25],[150,25],[200,50],[250,50],[300,50],[350,50],[400,50],[500,100],[600,100],[800,100],[1000,200],[1500,250],[2000,250]];
  for (const pair of brackets) if (target <= pair[0]) return pair;
  const yMax = Math.max(200, Math.ceil(target / 200) * 200);
  return [yMax, 200];
}
function lineValues(site, key) {
  if (key === "rx") return site.rx;
  const value = site[key];
  return site.rx.map(() => value);
}
function draw() {
  const site = VIEW.sites[current];
  const canvas = document.getElementById("chart");
  const stage = document.getElementById("stage");
  const dpr = window.devicePixelRatio || 1;
  const width = stage.clientWidth;
  const height = stage.clientHeight;
  canvas.width = Math.max(1, Math.floor(width * dpr));
  canvas.height = Math.max(1, Math.floor(height * dpr));
  canvas.style.width = width + "px";
  canvas.style.height = height + "px";
  const ctx = canvas.getContext("2d");
  ctx.setTransform(dpr, 0, 0, dpr, 0, 0);
  ctx.clearRect(0, 0, width, height);
  const pad = { l: 58, r: 16, t: 16, b: 72 };
  const plotW = width - pad.l - pad.r;
  const plotH = height - pad.t - pad.b;
  const packed = [];
  SERIES.forEach(spec => { if (on[spec.key]) packed.push(...lineValues(site, spec.key)); });
  const [yMax, yUnit] = yScale(packed);
  const n = VIEW.labels.length;
  const xAt = i => pad.l + (n <= 1 ? plotW / 2 : (i / (n - 1)) * plotW);
  const yAt = v => pad.t + plotH - (v / yMax) * plotH;
  ctx.strokeStyle = "#d9d9d9";
  ctx.fillStyle = "#595959";
  ctx.font = "12px Calibri, sans-serif";
  ctx.textAlign = "right";
  ctx.textBaseline = "middle";
  for (let y = 0; y <= yMax + 0.1; y += yUnit) {
    const py = yAt(y);
    ctx.beginPath();
    ctx.moveTo(pad.l, py);
    ctx.lineTo(width - pad.r, py);
    ctx.stroke();
    ctx.fillText(String(Math.round(y)), pad.l - 8, py);
  }
  const axisBottom = pad.t + plotH;
  ctx.strokeStyle = "#b0b0b0";
  ctx.beginPath();
  VIEW.labels.forEach((label, i) => {
    const x = xAt(i);
    ctx.moveTo(x, axisBottom);
    ctx.lineTo(x, axisBottom + 4);
  });
  ctx.stroke();
  ctx.fillStyle = "#595959";
  ctx.font = "8px Calibri, sans-serif";
  ctx.textAlign = "left";
  ctx.textBaseline = "middle";
  VIEW.labels.forEach((label, i) => {
    const time = label.slice(label.lastIndexOf(" ") + 1);
    ctx.save();
    ctx.translate(xAt(i), axisBottom + 40);
    ctx.rotate(-Math.PI / 2);
    ctx.fillText(time, 0, 0);
    ctx.restore();
  });
  ctx.font = "11px Calibri, sans-serif";
  ctx.textAlign = "center";
  ctx.textBaseline = "top";
  let dayStart = 0;
  VIEW.labels.forEach((label, i) => {
    const date = label.slice(0, label.lastIndexOf(" "));
    const next = VIEW.labels[i + 1];
    const nextDate = next ? next.slice(0, next.lastIndexOf(" ")) : "";
    if (date !== nextDate) {
      ctx.fillText(date, (xAt(dayStart) + xAt(i)) / 2, axisBottom + 46);
      dayStart = i + 1;
    }
  });
  ctx.strokeStyle = "#b0b0b0";
  ctx.beginPath();
  ctx.moveTo(pad.l, pad.t);
  ctx.lineTo(pad.l, pad.t + plotH);
  ctx.lineTo(width - pad.r, pad.t + plotH);
  ctx.stroke();
  SERIES.forEach(spec => {
    if (!on[spec.key]) return;
    const values = lineValues(site, spec.key);
    ctx.beginPath();
    ctx.strokeStyle = spec.color;
    ctx.lineWidth = spec.key === "rx" ? 2.25 : 1.5;
    ctx.setLineDash(spec.dash ? [6, 4] : []);
    let started = false;
    values.forEach((value, i) => {
      if (value == null || !isFinite(value)) { started = false; return; }
      const x = xAt(i), y = yAt(value);
      if (!started) { ctx.moveTo(x, y); started = true; }
      else ctx.lineTo(x, y);
    });
    ctx.stroke();
  });
  ctx.setLineDash([]);
  canvas._geom = { pad, plotW, plotH, n, xAt, yMax };
}
function showTip(event) {
  const canvas = document.getElementById("chart");
  const geom = canvas._geom;
  const tip = document.getElementById("tip");
  if (!geom) return;
  const rect = canvas.getBoundingClientRect();
  const x = event.clientX - rect.left;
  const y = event.clientY - rect.top;
  if (x < geom.pad.l || x > geom.pad.l + geom.plotW || y < geom.pad.t || y > geom.pad.t + geom.plotH) {
    tip.style.display = "none";
    return;
  }
  const i = Math.max(0, Math.min(geom.n - 1, Math.round((x - geom.pad.l) / geom.plotW * (geom.n - 1))));
  const site = VIEW.sites[current];
  const rows = SERIES.filter(spec => on[spec.key]).map(spec => {
    const value = lineValues(site, spec.key)[i];
    const shown = value == null ? "—" : value.toFixed(2);
    return `<div style="color:${spec.color}">${spec.name}: ${shown}</div>`;
  }).join("");
  tip.innerHTML = `<b>${VIEW.labels[i]}</b>${rows}`;
  tip.style.display = "block";
  const left = Math.min(x + 12, rect.width - 220);
  tip.style.left = Math.max(8, left) + "px";
  tip.style.top = Math.max(8, y + 12) + "px";
}
function legendName(spec, site) {
  const num = value => value == null || value === "" ? null : Number(value).toFixed(2);
  if (spec.key === "peak" && num(site.peak)) return `Peak (${num(site.peak)} Mbit/s)`;
  if (spec.key === "level" && num(site.level)) return `Below5%Peak (${num(site.level)} Mbit/s)`;
  if (spec.key === "license" && num(site.license)) return `License (${num(site.license)} Mbps)`;
  if (spec.key === "threshold" && num(site.threshold)) return `85% of license (${num(site.threshold)} Mbps)`;
  return spec.name;
}
function kpi(label, value, unit) {
  const shown = value == null || value === "" ? "—" : value;
  return `<div class="kpi"><b>${shown}${unit && shown !== "—" ? unit : ""}</b><span>${label}</span></div>`;
}
function renderChart() {
  const site = VIEW.sites[current];
  document.getElementById("title").textContent = site.listed && site.n
    ? `${String(site.n).padStart(2, "0")}    ${site.site}`
    : site.site;
  document.getElementById("caption").textContent = site.listed
    ? `${VIEW.period}    ·    ${site.category}    ·    ${site.severity}`
    : `${VIEW.period}    ·    Not listed`;
  document.getElementById("kpis").innerHTML = [
    kpi("Peak (Mbit/s)", site.peak, ""),
    kpi("Below5%Peak (Mbit/s)", site.level, ""),
    kpi("License (Mbps)", site.license, ""),
    kpi("85% of license", site.threshold, ""),
    kpi("Gap (%)", site.gap, ""),
    kpi("Peak ratio (%)", site.ratio, ""),
    kpi("Where touched", site.where, ""),
    kpi("Tech", site.tech, ""),
    kpi("Site Type", site.siteType, "")
  ].join("");
  const ossNames = ["DL Throughput 1h (Mbps)", "DLPRBUtilization,%", "Data Volume 1h (GB)", "Max User 1h", "Mean User 1h"];
  const ossPair = (VIEW.oss || {})[site.site.slice(0, 7)];
  function ossBlock(title, color, values) {
    if (!values) return `<section class="oss-block"><h2 style="background:${color}">${title}</h2><p class="oss-empty">No OSS row for this site code.</p></section>`;
    const rows = ossNames.map((name, i) => `<div class="oss-row"><span>${name}</span><b>${values[i] == null ? "—" : values[i]}</b></div>`).join("");
    return `<section class="oss-block"><h2 style="background:${color}">${title}</h2>${rows}</section>`;
  }
  document.querySelectorAll("#legend .leg-name").forEach(node => {
    const key = node.parentElement.querySelector("input").dataset.key;
    const spec = SERIES.find(item => item.key === key);
    node.textContent = legendName(spec, site);
  });
  document.getElementById("oss").innerHTML = [
    ossBlock("DataBusyHour (1 BusyHour)", "#6C3483", ossPair && ossPair[0]),
    ossBlock("UserBusyHour (1 BusyHour)", "#1A5276", ossPair && ossPair[1])
  ].join("");
  document.querySelectorAll(".site").forEach(button => {
    button.classList.toggle("on", Number(button.dataset.index) === current);
  });
  try { history.replaceState(null, "", "#" + encodeURIComponent(site.site)); } catch (err) {}
  draw();
}
function paintRows() {
  const spacer = document.getElementById("spacer");
  if (!spacer) return;
  const list = document.getElementById("list");
  if (!shown.length) {
    spacer.innerHTML = '<p class="sub" style="position:absolute;padding:12px">No sites match.</p>';
    return;
  }
  const start = Math.max(0, Math.floor(list.scrollTop / ROW_H) - 8);
  const end = Math.min(shown.length, Math.ceil((list.scrollTop + list.clientHeight) / ROW_H) + 8);
  spacer.innerHTML = shown.slice(start, end).map((site, offset) => `
    <button class="site ${site.i === current ? "on" : ""}" style="top:${(start + offset) * ROW_H}px" data-index="${site.i}" type="button">
      <div><span class="name">${rowTitle(site)}</span></div>
      <div class="meta">${rowMeta(site)}</div>
    </button>`).join("");
}
function renderList(scrollToSelection) {
  shown = visible();
  const list = document.getElementById("list");
  if (!document.getElementById("spacer")) list.innerHTML = '<div id="spacer"></div>';
  const spacer = document.getElementById("spacer");
  spacer.style.height = Math.max(shown.length, 1) * ROW_H + "px";
  if (scrollToSelection) {
    const pos = shown.findIndex(site => site.i === current);
    if (pos >= 0) list.scrollTop = Math.max(0, pos * ROW_H - 8);
  }
  paintRows();
  document.getElementById("filters").innerHTML = filters.map(name => {
    const count = name === "All Sites"
      ? VIEW.sites.length
      : name === "All"
        ? VIEW.sites.filter(site => site.listed).length
        : VIEW.sites.filter(site => site.listed && primary(site) === name).length;
    return `<button type="button" data-filter="${name}" class="${filter === name ? "on" : ""}">${name} ${count.toLocaleString()}</button>`;
  }).join("");
}
function selectIndex(index) {
  current = Math.max(0, Math.min(VIEW.sites.length - 1, index));
  renderList(true);
  renderChart();
}
function step(delta) {
  if (!shown.length) return;
  let pos = shown.findIndex(site => site.i === current);
  if (pos < 0) pos = 0;
  pos = Math.max(0, Math.min(shown.length - 1, pos + delta));
  selectIndex(shown[pos].i);
}
document.getElementById("list").addEventListener("click", event => {
  const button = event.target.closest(".site");
  if (!button) return;
  selectIndex(Number(button.dataset.index));
});
document.getElementById("filters").addEventListener("click", event => {
  const button = event.target.closest("button");
  if (!button) return;
  filter = button.dataset.filter;
  const rows = visible();
  if (rows.length && !rows.some(site => site.i === current)) current = rows[0].i;
  renderList(true);
  renderChart();
});
document.getElementById("list").addEventListener("scroll", paintRows);
document.getElementById("q").addEventListener("input", event => {
  query = event.target.value;
  const rows = visible();
  if (rows.length && !rows.some(site => site.i === current)) current = rows[0].i;
  renderList(true);
  renderChart();
});
document.getElementById("prev").addEventListener("click", () => step(-1));
document.getElementById("next").addEventListener("click", () => step(1));
document.getElementById("chart").addEventListener("mousemove", showTip);
document.getElementById("chart").addEventListener("mouseleave", () => {
  document.getElementById("tip").style.display = "none";
});
window.addEventListener("resize", draw);
  document.getElementById("legend").innerHTML = SERIES.map(spec => `
  <label><input type="checkbox" data-key="${spec.key}" checked>
  <i class="swatch ${spec.dash ? "dash" : ""}" style="border-top-color:${spec.color}"></i><span class="leg-name">${spec.name}</span></label>`).join("");
document.getElementById("legend").addEventListener("change", event => {
  const key = event.target.dataset.key;
  if (!key) return;
  on[key] = event.target.checked;
  draw();
});
const issueCount = VIEW.sites.filter(site => site.listed).length;
document.getElementById("period").textContent = `${VIEW.period} · ${issueCount.toLocaleString()} issue sites · All Sites ${VIEW.sites.length.toLocaleString()}`;
const hash = decodeURIComponent(location.hash.replace(/^#/, ""));
const hashed = VIEW.sites.findIndex(site => site.site === hash || String(site.n) === hash);
current = hashed >= 0 ? hashed : 0;
renderList(true);
renderChart();
</script>
</body>
</html>
"""


def write_html(view: dict, dest: Path) -> None:
    payload = json.dumps(view, separators=(",", ":")).replace("<", "\\u003c")
    page = (
        HTML.replace("__DATA__", payload)
        .replace("__EXCEL_VERSION__", EXCEL_VERSION)
    )
    dest.write_text(page, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook")
    parser.add_argument("-o", "--output", required=True)
    parser.add_argument("--feg", default="FEG_KPI_Eric.xlsb")
    parser.add_argument("--license", default="DL_License_File.xlsb")
    parser.add_argument("--geo", default="Physical_Site_Database_14Oct26.xlsx")
    args = parser.parse_args()
    view = load_chart_view(Path(args.workbook))
    view["htmlVersion"] = HTML_VERSION
    view["oss"] = load_oss_1h(Path(args.workbook))
    print(f"OSS 1h site codes: {len(view['oss']):,}", flush=True)
    feg_path = Path(args.feg)
    if feg_path.exists():
        add_all_sites(view, feg_path, Path(args.license), Path(args.geo))
    dest = Path(args.output)
    write_html(view, dest)
    hours = len(view["labels"])
    print(f"Wrote {dest} · {len(view['sites'])} sites · {hours} hours · {dest.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
